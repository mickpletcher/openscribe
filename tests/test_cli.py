from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import yaml
from docx import Document
from typer.testing import CliRunner

import openscribe.compile as compile_module
from openscribe.cli import app
from openscribe.project import list_chapters
from openscribe.proofreading import ProofreadingIssue, ProofreadingResult

runner = CliRunner()


def test_outline_shows_literal_status_text(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    result = runner.invoke(app, ["init", "North County"])
    assert result.exit_code == 0

    result = runner.invoke(app, ["new", "part", "Opening"])
    assert result.exit_code == 0

    result = runner.invoke(
        app,
        [
            "new",
            "chapter",
            "Arrival",
            "--part",
            "Opening",
            "--status",
            "draft",
        ],
    )
    assert result.exit_code == 0

    result = runner.invoke(app, ["outline"])
    assert result.exit_code == 0
    assert "Opening" in result.stdout
    assert "part-01-opening" not in result.stdout
    assert "Arrival [draft]" in result.stdout


def test_compile_writes_docx_output(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft"],
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli stepped off the bus into wet summer heat.\n",
        encoding="utf-8",
    )
    assert (
        runner.invoke(
            app,
            ["new", "scene", "Station", "--chapter", "Arrival", "--body", "The doors closed behind him."],
        ).exit_code
        == 0
    )

    result = runner.invoke(app, ["compile"])
    assert result.exit_code == 0

    output_path = tmp_path / "build" / "north-county.docx"
    assert output_path.exists()

    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert "North County" in paragraph_text
    assert "Opening" in paragraph_text
    assert "Arrival" in paragraph_text
    assert "Eli stepped off the bus into wet summer heat." in paragraph_text
    assert "The doors closed behind him." in paragraph_text
    assert not any("openscribe-scene-id" in paragraph for paragraph in paragraph_text)


def test_compile_writes_pdf_output(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft"],
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli stepped off the bus into wet summer heat.\n",
        encoding="utf-8",
    )

    result = runner.invoke(app, ["compile", "--format", "pdf"])
    assert result.exit_code == 0

    output_path = tmp_path / "build" / "north-county.pdf"
    assert output_path.exists()
    assert output_path.read_bytes().startswith(b"%PDF")


def test_compile_writes_epub_output(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft"],
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli stepped off the bus into wet summer heat.\n",
        encoding="utf-8",
    )

    result = runner.invoke(app, ["compile", "--format", "epub"])
    assert result.exit_code == 0

    output_path = tmp_path / "build" / "north-county.epub"
    assert output_path.exists()

    with ZipFile(output_path) as archive:
        names = set(archive.namelist())

    assert "mimetype" in names
    assert "EPUB/content.opf" in names
    assert "EPUB/nav.xhtml" in names


def test_compile_respects_project_compile_settings(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft"],
        ).exit_code
        == 0
    )

    config_path = tmp_path / ".openscribe" / "project.yaml"
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    config["compile"]["output_filename"] = "review-copy"
    config["compile"]["include_title_page"] = False
    config["compile"]["include_part_headings"] = False
    config["compile"]["chapter_heading_style"] = "chapter-number-title"
    config_path.write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli stepped off the bus into wet summer heat.\n",
        encoding="utf-8",
    )

    result = runner.invoke(app, ["compile"])
    assert result.exit_code == 0

    output_path = tmp_path / "build" / "review-copy.docx"
    assert output_path.exists()

    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert "North County" not in paragraph_text
    assert "Opening" not in paragraph_text
    assert "Chapter 1: Arrival" in paragraph_text


def test_compile_prefers_pandoc_when_available(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli stepped off the bus into wet summer heat.\n",
        encoding="utf-8",
    )

    command_log: list[list[str]] = []

    def fake_run(command: list[str], capture_output: bool, text: bool, check: bool):
        command_log.append(command)
        output_path = Path(command[command.index("--output") + 1])
        output_path.write_text("pandoc output", encoding="utf-8")

        class Result:
            returncode = 0
            stdout = ""
            stderr = ""

        return Result()

    monkeypatch.setattr(compile_module.shutil, "which", lambda name: "pandoc.exe" if name == "pandoc" else None)
    monkeypatch.setattr(compile_module.subprocess, "run", fake_run)

    result = runner.invoke(app, ["compile"])
    assert result.exit_code == 0

    output_path = tmp_path / "build" / "north-county.docx"
    assert output_path.exists()
    assert output_path.read_text(encoding="utf-8") == "pandoc output"
    assert command_log
    assert command_log[0][0] == "pandoc.exe"
    assert "--to" in command_log[0]
    assert "docx" in command_log[0]


def test_read_outputs_continuous_manuscript_using_template(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft"],
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8")
        + "Eli stepped off the bus into wet summer heat.\n\nHe kept walking toward Willow Lane.\n",
        encoding="utf-8",
    )

    result = runner.invoke(app, ["read", "--template", "manuscript"])
    assert result.exit_code == 0
    assert "North County" in result.stdout
    assert "Opening" not in result.stdout
    assert "Chapter 1: Arrival" in result.stdout
    assert "He kept walking toward Willow Lane." in result.stdout


def test_set_and_show_chapter_metadata(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    result = runner.invoke(
        app,
        [
            "set",
            "chapter",
            "Arrival",
            "--status",
            "revised",
            "--label",
            "action",
            "--pov",
            "Eli",
            "--word-target",
            "2500",
            "--synopsis",
            "Eli reaches town.",
            "--notes",
            "Tighten the station scene.",
        ],
    )
    assert result.exit_code == 0

    result = runner.invoke(app, ["show", "chapter", "Arrival"])
    assert result.exit_code == 0
    assert "status: revised" in result.stdout
    assert "label: action" in result.stdout
    assert "pov: Eli" in result.stdout
    assert "word_target: 2500" in result.stdout
    assert "synopsis: Eli reaches town." in result.stdout
    assert "notes: Tighten the station scene." in result.stdout


def test_set_and_show_part_metadata(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0

    result = runner.invoke(app, ["set", "part", "Opening", "--title", "Cold Open"])
    assert result.exit_code == 0

    result = runner.invoke(app, ["show", "part", "Cold Open"])
    assert result.exit_code == 0
    assert "title: Cold Open" in result.stdout


def test_story_idea_commands(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0

    result = runner.invoke(
        app,
        [
            "idea",
            "add",
            "The Flood Ledger",
            "--premise",
            "A county clerk finds a ledger that predicts deaths.",
            "--genre",
            "Southern Gothic",
            "--tone",
            "Uneasy",
            "--status",
            "seed",
            "--notes",
            "Tie the flood history to the missing records plot.",
        ],
    )
    assert result.exit_code == 0
    assert "notes\\story-ideas\\the-flood-ledger.md" in result.stdout

    result = runner.invoke(app, ["idea", "list"])
    assert result.exit_code == 0
    assert "The Flood" in result.stdout
    assert "Southern" in result.stdout
    assert "Ideas: 1" in result.stdout

    result = runner.invoke(app, ["show", "idea", "The Flood Ledger"])
    assert result.exit_code == 0
    assert "premise: A county clerk finds a ledger that predicts deaths." in result.stdout
    assert "genre: Southern Gothic" in result.stdout
    assert "tone: Uneasy" in result.stdout
    assert "status: seed" in result.stdout


def test_template_scene_index_and_snapshot_commands(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    result = runner.invoke(app, ["templates"])
    assert result.exit_code == 0
    assert "fiction" in result.stdout
    assert "research" in result.stdout
    assert "technical" in result.stdout
    assert "research-paper" in result.stdout

    assert runner.invoke(app, ["init", "North County", "--template", "technical"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    result = runner.invoke(app, ["new", "scene", "Bus Stop", "--chapter", "Arrival", "--body", "Eli arrives in town."])
    assert result.exit_code == 0

    result = runner.invoke(app, ["show", "chapter", "Arrival"])
    assert result.exit_code == 0
    assert "scene_count: 1" in result.stdout
    assert "Bus Stop" in result.stdout

    result = runner.invoke(app, ["index", "rebuild"])
    assert result.exit_code == 0
    assert ".openscribe\\index\\project-index.yaml" in result.stdout

    result = runner.invoke(app, ["index", "show"])
    assert result.exit_code == 0
    assert "chapter_count: 1" in result.stdout
    assert "scene_count: 1" in result.stdout

    result = runner.invoke(app, ["snapshot", "save", "before-rewrite"])
    assert result.exit_code == 0
    assert ".openscribe\\snapshots\\" in result.stdout

    result = runner.invoke(app, ["snapshot", "list"])
    assert result.exit_code == 0
    assert "before" in result.stdout
    assert "checkpoint" in result.stdout

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(chapter_path.read_text(encoding="utf-8") + "Changed later.\n", encoding="utf-8")

    result = runner.invoke(app, ["snapshot", "diff", "before-rewrite"])
    assert result.exit_code == 0
    assert "Changed later." in result.stdout

    result = runner.invoke(app, ["snapshot", "restore", "before-rewrite"])
    assert result.exit_code == 0
    assert "Preview only" in result.stdout
    assert "Changed later." in chapter_path.read_text(encoding="utf-8")

    result = runner.invoke(app, ["snapshot", "restore", "before-rewrite", "--apply"])
    assert result.exit_code == 0
    assert "Restored snapshot" in result.stdout
    assert "Automatic backup" in result.stdout
    assert "Changed later." not in chapter_path.read_text(encoding="utf-8")


def test_custom_template_import_and_workflow_commands(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "Template Source", "--template", "screenwriting"]).exit_code == 0

    result = runner.invoke(app, ["template", "save", "screenplay-custom"])
    assert result.exit_code == 0
    assert ".openscribe\\templates\\screenplay-custom.yaml" in result.stdout

    source_folder = tmp_path / "existing-manuscript"
    (source_folder / "act-one").mkdir(parents=True)
    (source_folder / "act-one" / "opening.md").write_text("# Opening\n\nImported text.\n", encoding="utf-8")

    target_folder = tmp_path / "imported-project"
    result = runner.invoke(
        app,
        [
            "import",
            "folder",
            str(source_folder),
            "--title",
            "Imported Screenplay",
            "--template-file",
            str(tmp_path / ".openscribe" / "templates" / "screenplay-custom.yaml"),
            "--path",
            str(target_folder),
        ],
    )
    assert result.exit_code == 0
    assert "Imported project to" in result.stdout

    monkeypatch.chdir(target_folder)
    assert runner.invoke(app, ["new", "part", "Second Act"]).exit_code == 0
    assert runner.invoke(app, ["workflow", "nonfiction-section", "Background", "--part", "Second Act"]).exit_code == 0
    result = runner.invoke(
        app,
        [
            "workflow",
            "screenplay-scene",
            "int. diner - night",
            "--chapter",
            "Background",
            "--body",
            "Two strangers wait.",
        ],
    )
    assert result.exit_code == 0

    result = runner.invoke(app, ["show", "chapter", "Background"])
    assert result.exit_code == 0
    assert "label: section" in result.stdout
    assert "INT. DINER - NIGHT" in result.stdout


def test_research_paper_and_conference_workflows(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "Grid Study", "--template", "research"]).exit_code == 0

    result = runner.invoke(app, ["workflow", "research-paper", "--part", "Paper", "--include-appendix"])
    assert result.exit_code == 0
    assert "Created research paper sections: 9" in result.stdout
    assert "ch-01-abstract.md" in result.stdout
    assert "ch-09-appendix.md" in result.stdout

    abstract_path = tmp_path / "manuscript" / "part-01-paper" / "ch-01-abstract.md"
    abstract_path.write_text(
        abstract_path.read_text(encoding="utf-8") + "This paper studies resilient rural energy systems.\n",
        encoding="utf-8",
    )
    intro_path = tmp_path / "manuscript" / "part-01-paper" / "ch-02-introduction.md"
    intro_path.write_text(
        intro_path.read_text(encoding="utf-8") + "Rural microgrids need better planning models.\n",
        encoding="utf-8",
    )

    compile_result = runner.invoke(app, ["compile", "--profile", "research-paper"])
    assert compile_result.exit_code == 0

    output_path = tmp_path / "build" / "grid-study-research-paper.docx"
    assert output_path.exists()
    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert "1. Abstract" in paragraph_text
    assert "2. Introduction" in paragraph_text

    result = runner.invoke(app, ["workflow", "conference-materials", "Grid Study 2026", "--venue", "EnergyConf"])
    assert result.exit_code == 0
    assert "Created conference materials: 9" in result.stdout
    assert "grid-study-2026-slide-draft.md" in result.stdout
    assert "grid-study-2026-submission-checklist.md" in result.stdout
    assert "grid-study-2026-timed-talk-plan.md" in result.stdout
    assert "grid-study-2026-submission-status.md" in result.stdout

    slide_path = tmp_path / "notes" / "presentations" / "grid-study-2026-slide-draft.md"
    checklist_path = tmp_path / "research" / "conferences" / "grid-study-2026-submission-checklist.md"
    timed_path = tmp_path / "notes" / "presentations" / "grid-study-2026-timed-talk-plan.md"
    status_path = tmp_path / "research" / "conferences" / "grid-study-2026-submission-status.md"
    assert slide_path.exists()
    assert checklist_path.exists()
    assert timed_path.exists()
    assert status_path.exists()
    assert "Title Slide" in slide_path.read_text(encoding="utf-8")
    assert "EnergyConf" in checklist_path.read_text(encoding="utf-8")
    assert "Run Of Show" in timed_path.read_text(encoding="utf-8")
    assert "Current State" in status_path.read_text(encoding="utf-8")

    schedule_path = tmp_path / "energyconf-schedule.csv"
    schedule_path.write_text(
        "Title,Day,Time,Room,Format,Presenter,Status,Notes\n"
        "Microgrid Timing,Day 1,09:00,Hall A,Talk,Eli Harper,accepted,Need a shorter opening.\n"
        "Poster Review,Day 2,14:00,Hall B,Poster,Nora Bell,planned,Confirm print size.\n",
        encoding="utf-8",
    )

    result = runner.invoke(
        app,
        [
            "workflow",
            "conference-schedule-import",
            str(schedule_path),
            "--venue",
            "EnergyConf",
        ],
    )
    assert result.exit_code == 0
    assert "Imported conference schedule: 4" in result.stdout
    assert "energyconf-session-schedule.md" in result.stdout
    assert "energyconf-session-checklist.md" in result.stdout
    assert "energyconf-01-microgrid-timing.md" in result.stdout

    schedule_output = tmp_path / "research" / "conferences" / "energyconf-session-schedule.md"
    session_checklist_path = tmp_path / "research" / "conferences" / "energyconf-session-checklist.md"
    session_note_path = tmp_path / "research" / "conferences" / "sessions" / "energyconf-01-microgrid-timing.md"
    assert schedule_output.exists()
    assert session_checklist_path.exists()
    assert session_note_path.exists()
    assert "Microgrid Timing" in schedule_output.read_text(encoding="utf-8")
    assert "Poster Review" in session_checklist_path.read_text(encoding="utf-8")
    assert "Session Checklist" in session_note_path.read_text(encoding="utf-8")


def test_nonfiction_source_and_citation_workflows(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "County History", "--template", "nonfiction"]).exit_code == 0

    result = runner.invoke(
        app,
        [
            "workflow",
            "source-note",
            "River Ledger Study",
            "--type",
            "article",
            "--author",
            "J. Harper",
            "--year",
            "2024",
            "--url",
            "https://example.com/ledger",
            "--notes",
            "Supports the flood records argument.",
        ],
    )
    assert result.exit_code == 0
    assert "research\\sources\\river-ledger-study.md" in result.stdout

    result = runner.invoke(app, ["workflow", "citation-pack", "--style", "Chicago"])
    assert result.exit_code == 0
    assert "Created citation tracking files:" in result.stdout
    assert "research\\source-usage-map.md" in result.stdout

    source_path = tmp_path / "research" / "sources" / "river-ledger-study.md"
    citation_log_path = tmp_path / "research" / "citation-log.md"
    bibliography_path = tmp_path / "research" / "bibliography-notes.md"
    assert source_path.exists()
    assert citation_log_path.exists()
    assert bibliography_path.exists()
    source_text = source_path.read_text(encoding="utf-8")
    assert "title: River Ledger Study" in source_text
    assert "author: J. Harper" in source_text
    assert "Supports the flood records argument." in source_text
    assert "| Section | Source | Use | Notes |" in citation_log_path.read_text(encoding="utf-8")
    assert "## Style" in bibliography_path.read_text(encoding="utf-8")


def test_board_visual_commands(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["board", "note", "add", "Clue", "--group", "plot"]).exit_code == 0
    assert runner.invoke(app, ["board", "note", "add", "Threat", "--group", "plot"]).exit_code == 0

    result = runner.invoke(app, ["board", "note", "move", "note-001", "--x", "10", "--y", "3"])
    assert result.exit_code == 0
    assert "Moved note-001" in result.stdout

    result = runner.invoke(app, ["board", "layout", "auto"])
    assert result.exit_code == 0
    assert "Applied board auto layout" in result.stdout

    result = runner.invoke(app, ["board", "view", "--width", "50", "--height", "10"])
    assert result.exit_code == 0
    assert "Board View" in result.stdout


def test_find_chapters_filters_by_metadata_and_text(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Ending"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            [
                "new",
                "chapter",
                "Arrival",
                "--part",
                "Opening",
                "--status",
                "draft",
                "--label",
                "setup",
                "--pov",
                "Eli",
                "--synopsis",
                "Eli reaches town.",
            ],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            [
                "new",
                "chapter",
                "Departure",
                "--part",
                "Ending",
                "--status",
                "revised",
                "--label",
                "finale",
                "--pov",
                "Nora",
                "--notes",
                "Need stronger station callback.",
            ],
        ).exit_code
        == 0
    )

    arrival_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    arrival_path.write_text(
        arrival_path.read_text(encoding="utf-8") + "The station was almost empty.\n",
        encoding="utf-8",
    )
    departure_path = tmp_path / "manuscript" / "part-02-ending" / "ch-01-departure.md"
    departure_path.write_text(
        departure_path.read_text(encoding="utf-8") + "Nora left before sunrise.\n",
        encoding="utf-8",
    )

    result = runner.invoke(app, ["find", "chapters", "--status", "revised"])
    assert result.exit_code == 0
    assert "Departure" in result.stdout
    assert "Arrival" not in result.stdout
    assert "Matches: 1" in result.stdout

    result = runner.invoke(app, ["find", "chapters", "--part", "Opening", "--pov", "Eli"])
    assert result.exit_code == 0
    assert "Arrival" in result.stdout
    assert "Departure" not in result.stdout

    result = runner.invoke(app, ["find", "chapters", "--text", "station"])
    assert result.exit_code == 0
    assert "Arrival" in result.stdout
    assert "Departure" in result.stdout
    assert "Matches: 2" in result.stdout

    assert (
        runner.invoke(
            app,
            ["new", "scene", "Station Watch", "--chapter", "Arrival", "--body", "Scene line about the station bell."],
        ).exit_code
        == 0
    )
    result = runner.invoke(app, ["find", "scenes", "--text", "station bell"])
    assert result.exit_code == 0
    assert "Station Watch" in result.stdout
    assert "Arrival" in result.stdout


def test_report_project_and_batch_update(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Ending"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft", "--label", "setup", "--pov", "Eli"],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            [
                "new",
                "chapter",
                "Departure",
                "--part",
                "Ending",
                "--status",
                "draft",
                "--label",
                "finale",
                "--pov",
                "Nora",
            ],
        ).exit_code
        == 0
    )

    result = runner.invoke(app, ["report", "project"])
    assert result.exit_code == 0
    assert "By Status" in result.stdout
    assert "draft" in result.stdout
    assert "By POV" in result.stdout
    assert "Eli" in result.stdout

    result = runner.invoke(app, ["set", "chapters", "--match-status", "draft", "--status", "revised"])
    assert result.exit_code == 0
    assert "Updated chapters: 2" in result.stdout

    result = runner.invoke(app, ["find", "chapters", "--status", "revised"])
    assert result.exit_code == 0
    assert "Matches: 2" in result.stdout

    assert (
        runner.invoke(
            app, ["new", "scene", "Open Road", "--chapter", "Arrival", "--body", "Eli starts the drive."]
        ).exit_code
        == 0
    )
    result = runner.invoke(app, ["report", "scenes"])
    assert result.exit_code == 0
    assert "Scenes By Chapter" in result.stdout
    assert "Arrival" in result.stdout


def test_board_note_workflow_and_promote(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0

    result = runner.invoke(
        app,
        [
            "board",
            "note",
            "add",
            "Station Secret",
            "--body",
            "The station master is hiding records.",
            "--group",
            "plot",
        ],
    )
    assert result.exit_code == 0
    assert "note-001" in result.stdout

    result = runner.invoke(app, ["board", "note", "add", "Eli Clue", "--body", "Eli finds a ledger."])
    assert result.exit_code == 0
    assert "note-002" in result.stdout

    assert runner.invoke(app, ["board", "link", "add", "note-001", "note-002"]).exit_code == 0
    assert runner.invoke(app, ["board", "group", "set", "note-002", "plot"]).exit_code == 0

    result = runner.invoke(app, ["board", "note", "list", "--group", "plot"])
    assert result.exit_code == 0
    assert "Station Secret" in result.stdout
    assert "Eli Clue" in result.stdout

    result = runner.invoke(app, ["board", "promote", "note-001", "--part", "Opening", "--chapter", "Station Secret"])
    assert result.exit_code == 0
    assert "Station Secret" in result.stdout

    result = runner.invoke(app, ["find", "chapters", "--text", "hiding records"])
    assert result.exit_code == 0
    assert "Station Secret" in result.stdout

    result = runner.invoke(app, ["board", "chapter", "add", "note-002", "Station Secret"])
    assert result.exit_code == 0
    assert "Linked note-002 to chapter Station Secret" in result.stdout

    result = runner.invoke(app, ["board", "note", "list"])
    assert result.exit_code == 0
    board_text = (tmp_path / ".openscribe" / "boards" / "default.yaml").read_text(encoding="utf-8")
    chapter_id = next(chapter.chapter_id for chapter in list_chapters(tmp_path) if chapter.title == "Station Secret")
    assert chapter_id in board_text

    result = runner.invoke(app, ["board", "chapter", "remove", "note-002", "Station Secret"])
    assert result.exit_code == 0
    assert "Removed chapter link" in result.stdout


def test_element_workflow_and_appears_in(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Marcus Vale met the station keeper at dawn.\n",
        encoding="utf-8",
    )

    result = runner.invoke(
        app, ["element", "add", "character", "Marcus Vale", "--notes", "Main investigator", "--tags", "lead,viewpoint"]
    )
    assert result.exit_code == 0
    assert "cha-marcus-vale" in result.stdout

    assert runner.invoke(app, ["element", "alias", "add", "cha-marcus-vale", "Marcus"]).exit_code == 0
    assert runner.invoke(app, ["element", "add", "setting", "North Station"]).exit_code == 0
    assert (
        runner.invoke(app, ["element", "relate", "cha-marcus-vale", "set-north-station", "--type", "visits"]).exit_code
        == 0
    )

    result = runner.invoke(app, ["element", "show", "cha-marcus-vale"])
    assert result.exit_code == 0
    assert "Marcus Vale" in result.stdout
    assert "Main investigator" in result.stdout
    assert "visits" in result.stdout

    result = runner.invoke(app, ["element", "appears-in", "Marcus"])
    assert result.exit_code == 0
    assert "Arrival" in result.stdout
    assert "Matches: 1" in result.stdout


def test_element_batch_updates_and_relation_removal(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["element", "add", "character", "Eli Harper", "--tags", "lead"]).exit_code == 0
    assert runner.invoke(app, ["element", "add", "setting", "North Station", "--tags", "place"]).exit_code == 0
    assert (
        runner.invoke(app, ["element", "relate", "cha-eli-harper", "set-north-station", "--type", "visits"]).exit_code
        == 0
    )

    result = runner.invoke(
        app,
        ["element", "set-many", "--match-type", "character", "--add-tags", "viewpoint", "--notes", "Lead investigator"],
    )
    assert result.exit_code == 0
    assert "Updated elements: 1" in result.stdout

    result = runner.invoke(app, ["element", "unrelate", "cha-eli-harper", "set-north-station", "--type", "visits"])
    assert result.exit_code == 0
    assert "Removed relation" in result.stdout

    result = runner.invoke(app, ["element", "show", "cha-eli-harper"])
    assert result.exit_code == 0
    assert "viewpoint" in result.stdout
    assert "Lead investigator" in result.stdout
    assert "visits" not in result.stdout


def test_multi_part_read_and_export_regression(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Ending"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Reckoning", "--part", "Ending"]).exit_code == 0

    arrival_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    arrival_path.write_text(arrival_path.read_text(encoding="utf-8") + "Opening text.\n", encoding="utf-8")
    reckoning_path = tmp_path / "manuscript" / "part-02-ending" / "ch-01-reckoning.md"
    reckoning_path.write_text(reckoning_path.read_text(encoding="utf-8") + "Ending text.\n", encoding="utf-8")

    read_result = runner.invoke(app, ["read"])
    assert read_result.exit_code == 0
    assert read_result.stdout.index("Opening") < read_result.stdout.index("Arrival")
    assert read_result.stdout.index("Ending") < read_result.stdout.index("Reckoning")
    assert read_result.stdout.index("Arrival") < read_result.stdout.index("Reckoning")

    compile_result = runner.invoke(app, ["compile"])
    assert compile_result.exit_code == 0

    output_path = tmp_path / "build" / "north-county.docx"
    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert paragraph_text.index("Opening") < paragraph_text.index("Arrival")
    assert paragraph_text.index("Ending") < paragraph_text.index("Reckoning")


def test_outliner_combines_structure_metadata_and_word_counts(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            [
                "new",
                "chapter",
                "Arrival",
                "--part",
                "Opening",
                "--status",
                "draft",
                "--label",
                "setup",
                "--pov",
                "Eli",
                "--word-target",
                "1200",
                "--synopsis",
                "Eli reaches town.",
            ],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app, ["new", "scene", "Cold Open", "--chapter", "Arrival", "--body", "Rain on the station."]
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli walked into town through the rain.\n",
        encoding="utf-8",
    )

    result = runner.invoke(app, ["outliner"])
    assert result.exit_code == 0
    assert "Project: North County" in result.stdout
    assert "Opening" in result.stdout
    assert "Arrival | status=draft | label=setup | pov=Eli" in result.stdout
    assert "target=1200" in result.stdout
    assert "scenes=1" in result.stdout
    assert "Synopsis: Eli reaches town." in result.stdout
    assert "Scenes: Cold Open" in result.stdout
    assert "Part Total Words:" in result.stdout


def test_move_part_and_chapter_reorders_manuscript(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Middle"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Crossing", "--part", "Middle"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Signal", "--part", "Middle"]).exit_code == 0

    result = runner.invoke(app, ["move", "part", "Middle", "--position", "1"])
    assert result.exit_code == 0
    assert "Moved part to position 1" in result.stdout

    result = runner.invoke(app, ["move", "chapter", "Signal", "--position", "1", "--part", "Opening"])
    assert result.exit_code == 0
    assert "Moved chapter to position 1" in result.stdout

    outline_result = runner.invoke(app, ["outline"])
    assert outline_result.exit_code == 0
    assert outline_result.stdout.index("Middle") < outline_result.stdout.index("Opening")
    assert outline_result.stdout.index("Signal") < outline_result.stdout.index("Arrival")

    moved_path = tmp_path / "manuscript" / "part-02-opening" / "ch-01-signal.md"
    assert moved_path.exists()


def test_compile_profiles_apply_expected_defaults(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(compile_module.shutil, "which", lambda name: None)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Eli stepped off the bus into wet summer heat.\n",
        encoding="utf-8",
    )

    print_result = runner.invoke(app, ["compile", "--profile", "print"])
    assert print_result.exit_code == 0

    print_output = tmp_path / "build" / "north-county-print.docx"
    assert print_output.exists()
    print_document = Document(print_output)
    print_paragraphs = [paragraph.text for paragraph in print_document.paragraphs if paragraph.text]
    assert "Opening" in print_paragraphs
    assert "Arrival" in print_paragraphs

    submission_result = runner.invoke(app, ["compile", "--profile", "submission"])
    assert submission_result.exit_code == 0

    submission_output = tmp_path / "build" / "north-county-submission.docx"
    assert submission_output.exists()
    submission_document = Document(submission_output)
    submission_paragraphs = [paragraph.text for paragraph in submission_document.paragraphs if paragraph.text]
    assert "Opening" not in submission_paragraphs
    assert "Chapter 1: Arrival" in submission_paragraphs

    ebook_result = runner.invoke(app, ["compile", "--profile", "ebook"])
    assert ebook_result.exit_code == 0
    ebook_output = tmp_path / "build" / "north-county-ebook.epub"
    assert ebook_output.exists()


def test_research_compile_settings_add_bibliography(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(compile_module.shutil, "which", lambda name: None)

    assert runner.invoke(app, ["init", "County Paper", "--template", "research"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Paper"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Introduction", "--part", "Paper"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            [
                "workflow",
                "source-note",
                "County Archive",
                "--author",
                "Stewart County",
                "--year",
                "1987",
                "--url",
                "https://example.com/archive",
            ],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            [
                "set",
                "compile-research",
                "--citation-style",
                "Chicago",
                "--include-bibliography",
                "--bibliography-title",
                "Works Cited",
            ],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            ["set", "goals", "--draft-word-target", "5000", "--session-word-target", "750", "--deadline", "2026-07-01"],
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-paper" / "ch-01-introduction.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "County Archive supports the opening claim.\n", encoding="utf-8"
    )

    result = runner.invoke(app, ["compile", "--profile", "research-paper"])
    assert result.exit_code == 0

    output_path = tmp_path / "build" / "county-paper-research-paper.docx"
    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert "Works Cited" in paragraph_text
    assert any("County Archive" in paragraph for paragraph in paragraph_text)

    report_result = runner.invoke(app, ["report", "project"])
    assert report_result.exit_code == 0
    assert "Draft target: 5000" in report_result.stdout
    assert "Deadline: 2026-07-01" in report_result.stdout

    cite_result = runner.invoke(
        app, ["workflow", "cite", "--chapter", "Introduction", "--source", "County Archive", "--style", "Chicago"]
    )
    assert cite_result.exit_code == 0
    chapter_text = chapter_path.read_text(encoding="utf-8")
    assert "Citation:" in chapter_text
    assert "Stewart County" in chapter_text


def test_open_helpers_use_editor_targets(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    opened: list[str] = []
    monkeypatch.setattr("openscribe.project.os.startfile", lambda path: opened.append(path))

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    result = runner.invoke(app, ["open", "chapter", "Arrival"])
    assert result.exit_code == 0
    assert any(path.endswith("ch-01-arrival.md") for path in opened)

    result = runner.invoke(app, ["open", "part", "Opening"])
    assert result.exit_code == 0
    assert any(path.endswith("part-01-opening") for path in opened)

    result = runner.invoke(app, ["open", "search", "--text", "Arrival"])
    assert result.exit_code == 0
    assert len(opened) == 3


def test_outliner_filters_and_tui_quick_actions(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)

    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening", "--status", "draft", "--label", "setup", "--pov", "Eli"],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            [
                "new",
                "chapter",
                "Departure",
                "--part",
                "Opening",
                "--status",
                "revised",
                "--label",
                "finale",
                "--pov",
                "Nora",
            ],
        ).exit_code
        == 0
    )

    result = runner.invoke(app, ["outliner", "--status", "revised"])
    assert result.exit_code == 0
    assert "Departure" in result.stdout
    assert "Arrival" not in result.stdout


def test_hosted_ai_requires_and_displays_data_transfer_approval(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "chapter", "Arrival", "--part", "Opening"],
        ).exit_code
        == 0
    )

    config_path = tmp_path / ".openscribe" / "project.yaml"
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    config["ai"]["enabled"] = True
    config_path.write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Private manuscript text.\n",
        encoding="utf-8",
    )

    denied = runner.invoke(app, ["ai", "summarize", "Arrival"])
    assert denied.exit_code == 2
    assert "complete chapter text" in denied.stderr

    approved = runner.invoke(
        app,
        ["ai", "summarize", "Arrival", "--allow-data-transfer"],
    )
    assert approved.exit_code == 2
    assert "Sending" in approved.stdout
    assert "hosted provider 'openai'" in approved.stdout
    assert "OPENAI_API_KEY is not set" in approved.stderr


def test_scoped_ai_commands_are_read_only_and_route_expected_tasks(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0
    config_path = tmp_path / ".openscribe" / "project.yaml"
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    config["ai"]["enabled"] = True
    config_path.write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")
    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Private manuscript text.\n",
        encoding="utf-8",
    )
    before = chapter_path.read_bytes()
    calls: list[tuple[str, str]] = []

    def fake_task(text, settings, context_label, task, **kwargs):
        calls.append((task, kwargs.get("question", "")))
        return f"result for {task}"

    monkeypatch.setattr("openscribe.cli.run_ai_task", fake_task)
    commands = [
        ["rewrite", "Arrival"],
        ["outline", "Arrival"],
        ["analyze", "Arrival", "--focus", "pacing"],
        ["analyze", "Arrival", "--focus", "continuity"],
        ["analyze", "Arrival", "--focus", "pov"],
        ["analyze", "Arrival", "--focus", "prose"],
        ["metadata", "Arrival"],
        ["brainstorm", "Arrival", "--question", "What breaks next?"],
        ["query", "Who arrived?"],
    ]
    for command in commands:
        result = runner.invoke(app, ["ai", *command, "--allow-data-transfer"])
        assert result.exit_code == 0, result.output
        assert "result for" in result.stdout

    assert [task for task, _ in calls] == [
        "rewrite",
        "outline",
        "pacing",
        "continuity",
        "point-of-view",
        "prose",
        "metadata",
        "brainstorm",
        "query",
    ]
    assert calls[-2][1] == "What breaks next?"
    assert calls[-1][1] == "Who arrived?"
    assert chapter_path.read_bytes() == before


def test_ai_analyze_rejects_unknown_focus(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0

    result = runner.invoke(app, ["ai", "analyze", "missing", "--focus", "sentiment"])

    assert result.exit_code == 2
    assert "Focus must be" in result.stderr


def test_proofread_chapter_renders_languagetool_findings(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    config_path = tmp_path / ".openscribe" / "project.yaml"
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    config["proofreading"]["enabled"] = True
    config_path.write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")
    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "This are wrong.\n",
        encoding="utf-8",
    )

    captured = {}

    def fake_check_text(text, settings, *, language, allow_data_transfer):
        captured["text"] = text
        captured["language"] = language
        captured["allow_data_transfer"] = allow_data_transfer
        return ProofreadingResult(
            issues=(
                ProofreadingIssue(
                    message="Possible agreement error.",
                    short_message="Agreement",
                    offset=5,
                    length=3,
                    replacements=("is",),
                    rule_id="THIS_NNS",
                    category="Grammar",
                    issue_type="grammar",
                    context="This are wrong.",
                ),
            ),
            language="en-US",
            software_version="6.6",
            incomplete_results=False,
        )

    monkeypatch.setattr("openscribe.cli.check_text", fake_check_text)

    result = runner.invoke(app, ["proofread", "chapter", "Arrival", "--language", "en-GB"])

    assert result.exit_code == 0
    assert captured == {
        "text": "This are wrong.\n",
        "language": "en-GB",
        "allow_data_transfer": False,
    }
    assert "LanguageTool: Arrival" in result.stdout
    assert "1:6" in result.stdout
    assert "THIS_NNS" in result.stdout
    assert "Findings: 1" in result.stdout


def test_hosted_proofreading_requires_and_displays_data_transfer_approval(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.chdir(tmp_path)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0

    config_path = tmp_path / ".openscribe" / "project.yaml"
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    config["proofreading"].update(
        {
            "enabled": True,
            "endpoint": "https://proofreading.example.test/v2/check",
        }
    )
    config_path.write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")
    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    chapter_path.write_text(
        chapter_path.read_text(encoding="utf-8") + "Private manuscript text.\n",
        encoding="utf-8",
    )

    denied = runner.invoke(app, ["proofread", "chapter", "Arrival"])
    assert denied.exit_code == 2
    assert "complete chapter text" in denied.stderr

    def fake_check_text(text, settings, *, language, allow_data_transfer):
        assert text == "Private manuscript text.\n"
        assert allow_data_transfer is True
        return ProofreadingResult((), "en-US", "enterprise", False)

    monkeypatch.setattr("openscribe.cli.check_text", fake_check_text)
    approved = runner.invoke(
        app,
        ["proofread", "chapter", "Arrival", "--allow-data-transfer"],
    )

    assert approved.exit_code == 0
    assert "Sending" in approved.stdout
    assert "hosted LanguageTool endpoint" in approved.stdout
    assert "No findings" in approved.stdout


def test_scene_commands_preview_then_apply(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0
    assert runner.invoke(app, ["new", "part", "Opening"]).exit_code == 0
    assert runner.invoke(app, ["new", "chapter", "Arrival", "--part", "Opening"]).exit_code == 0
    assert (
        runner.invoke(
            app,
            ["new", "scene", "Bus Stop", "--chapter", "Arrival", "--body", "First half. Second half."],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            ["new", "scene", "Town Hall", "--chapter", "Arrival", "--body", "Questions begin."],
        ).exit_code
        == 0
    )

    chapter_path = tmp_path / "manuscript" / "part-01-opening" / "ch-01-arrival.md"
    before = chapter_path.read_text(encoding="utf-8")
    preview = runner.invoke(
        app,
        ["scene", "move", "Bus Stop", "--chapter", "Arrival", "--position", "2"],
    )
    assert preview.exit_code == 0
    assert "Preview only" in preview.stdout
    assert chapter_path.read_text(encoding="utf-8") == before

    applied = runner.invoke(
        app,
        ["scene", "move", "Bus Stop", "--chapter", "Arrival", "--position", "2", "--apply"],
    )
    assert applied.exit_code == 0
    assert "Applied scene operation" in applied.stdout
    assert chapter_path.read_text(encoding="utf-8").index("## Town Hall") < chapter_path.read_text(
        encoding="utf-8"
    ).index("## Bus Stop")

    split = runner.invoke(
        app,
        [
            "scene",
            "split",
            "Bus Stop",
            "--chapter",
            "Arrival",
            "--at-text",
            "Second half",
            "--new-title",
            "After the Pause",
            "--apply",
        ],
    )
    assert split.exit_code == 0
    assert "## After the Pause" in chapter_path.read_text(encoding="utf-8")

    merge = runner.invoke(
        app,
        [
            "scene",
            "merge",
            "Bus Stop",
            "--with",
            "After the Pause",
            "--chapter",
            "Arrival",
            "--apply",
        ],
    )
    assert merge.exit_code == 0
    assert "## After the Pause" not in chapter_path.read_text(encoding="utf-8")


def test_migrate_commands_report_current_project_version(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert runner.invoke(app, ["init", "North County"]).exit_code == 0

    status = runner.invoke(app, ["migrate", "status"])
    assert status.exit_code == 0
    assert "Project format: 3" in status.stdout
    assert "No migration is required" in status.stdout

    apply = runner.invoke(app, ["migrate", "apply"])
    assert apply.exit_code == 0
    assert "No migration is required" in apply.stdout
