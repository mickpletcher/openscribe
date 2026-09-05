from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from html import escape
from pathlib import Path

from docx import Document
from ebooklib import epub
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from openscribe.project import (
    ChapterDocument,
    list_chapters,
    list_source_notes,
    load_project_config,
    slugify,
    strip_scene_markers,
)


class CompileError(RuntimeError):
    pass


@dataclass(slots=True)
class CompileOptions:
    format_name: str
    profile_name: str
    template_name: str
    output_filename: str
    include_title_page: bool
    include_part_headings: bool
    chapter_heading_style: str
    backend_name: str
    citation_style: str
    include_bibliography: bool
    include_reference_heading: bool
    bibliography_title: str


TEMPLATE_PRESETS = {
    "novel": {
        "include_title_page": True,
        "include_part_headings": True,
        "chapter_heading_style": "title-only",
    },
    "manuscript": {
        "include_title_page": True,
        "include_part_headings": False,
        "chapter_heading_style": "chapter-number-title",
    },
    "minimal": {
        "include_title_page": False,
        "include_part_headings": False,
        "chapter_heading_style": "title-only",
    },
    "academic": {
        "include_title_page": True,
        "include_part_headings": False,
        "chapter_heading_style": "section-number-title",
    },
}


PROFILE_PRESETS = {
    "print": {
        "format_name": "docx",
        "template_name": "novel",
        "include_title_page": True,
        "include_part_headings": True,
        "chapter_heading_style": "title-only",
        "output_suffix": "print",
    },
    "ebook": {
        "format_name": "epub",
        "template_name": "novel",
        "include_title_page": True,
        "include_part_headings": True,
        "chapter_heading_style": "title-only",
        "output_suffix": "ebook",
    },
    "submission": {
        "format_name": "docx",
        "template_name": "manuscript",
        "include_title_page": True,
        "include_part_headings": False,
        "chapter_heading_style": "chapter-number-title",
        "output_suffix": "submission",
    },
    "research-paper": {
        "format_name": "docx",
        "template_name": "academic",
        "include_title_page": True,
        "include_part_headings": False,
        "chapter_heading_style": "section-number-title",
        "output_suffix": "research-paper",
    },
}


def compile_project(
    root: Path,
    format_name: str | None = None,
    profile_name: str | None = None,
    template_name: str | None = None,
    output_path: Path | None = None,
) -> Path:
    config = load_project_config(root)
    project_title = str(config.get("title", "Untitled Project"))
    options = _resolve_compile_options(config, format_name, profile_name, template_name, output_path)
    chapters = _compiled_chapters(root)
    target_path = output_path or default_output_path(root, project_title, options)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    if _can_use_pandoc(options):
        _compile_project_with_pandoc(
            root,
            project_title,
            str(config.get("author", "")).strip(),
            chapters,
            target_path,
            options,
        )
        return target_path

    if options.format_name == "docx":
        _compile_project_to_docx(
            root,
            project_title,
            str(config.get("author", "")).strip(),
            chapters,
            target_path,
            options,
        )
        return target_path
    if options.format_name == "pdf":
        _compile_project_to_pdf(
            root,
            project_title,
            str(config.get("author", "")).strip(),
            chapters,
            target_path,
            options,
        )
        return target_path
    if options.format_name == "epub":
        _compile_project_to_epub(
            root,
            project_title,
            str(config.get("author", "")).strip(),
            chapters,
            target_path,
            options,
        )
        return target_path

    raise CompileError(f"Format '{options.format_name}' is not supported yet. Use docx, pdf, or epub.")


def assemble_manuscript_text(root: Path, template_name: str | None = None) -> str:
    config = load_project_config(root)
    project_title = str(config.get("title", "Untitled Project"))
    author = str(config.get("author", "")).strip()
    options = _resolve_compile_options(config, None, None, template_name, None)
    chapters = _compiled_chapters(root)

    sections: list[str] = []
    if options.include_title_page:
        title_lines = [project_title]
        if author:
            title_lines.append(author)
        sections.append("\n".join(title_lines))

    current_part = None
    for chapter_number, chapter in enumerate(chapters, start=1):
        if chapter.part_id != current_part:
            current_part = chapter.part_id
            if options.include_part_headings:
                sections.append(chapter.part)

        chapter_heading = _chapter_heading(chapter.title, chapter_number, options)
        chapter_sections = [chapter_heading, strip_scene_markers(chapter.body).strip()]
        sections.append("\n\n".join([part for part in chapter_sections if part]))

    bibliography_sections = _bibliography_sections(root, options)
    if bibliography_sections:
        sections.extend(bibliography_sections)

    return "\n\n".join(section for section in sections if section.strip())


def default_output_path(root: Path, project_title: str, options: CompileOptions) -> Path:
    if options.output_filename.strip():
        filename = options.output_filename.strip()
    else:
        filename = slugify(project_title)
        if options.profile_name:
            filename = f"{filename}-{options.profile_name}"

    if not filename.lower().endswith(f".{options.format_name}"):
        filename = f"{filename}.{options.format_name}"
    return root / "build" / filename


def _resolve_format(config: dict, format_name: str | None, output_path: Path | None, profile_name: str | None) -> str:
    if format_name:
        return format_name.strip().lower()
    if output_path and output_path.suffix:
        return output_path.suffix.lstrip(".").lower()
    if profile_name:
        return str(PROFILE_PRESETS[profile_name]["format_name"]).lower()
    compile_config = config.get("compile", {})
    return str(compile_config.get("default_format", "docx")).lower()


def _resolve_compile_options(
    config: dict,
    format_name: str | None,
    profile_name: str | None,
    template_name: str | None,
    output_path: Path | None,
) -> CompileOptions:
    compile_config = config.get("compile", {})
    research_compile = compile_config.get("research", {})
    resolved_profile = (profile_name or str(compile_config.get("default_profile", "") or "")).strip().lower()
    if resolved_profile and resolved_profile not in PROFILE_PRESETS:
        supported_profiles = ", ".join(sorted(PROFILE_PRESETS))
        raise CompileError(f"Profile '{resolved_profile}' is not supported yet. Use {supported_profiles}.")

    profile_defaults = PROFILE_PRESETS.get(resolved_profile, {})
    explicit_template = template_name is not None
    resolved_template = (
        (
            template_name
            or str(profile_defaults.get("template_name", ""))
            or str(compile_config.get("default_template", "novel"))
        )
        .strip()
        .lower()
    )
    if resolved_template not in TEMPLATE_PRESETS:
        supported = ", ".join(sorted(TEMPLATE_PRESETS))
        raise CompileError(f"Template '{resolved_template}' is not supported yet. Use {supported}.")

    template_defaults = TEMPLATE_PRESETS[resolved_template]
    resolved_format = _resolve_format(config, format_name, output_path, resolved_profile or None)
    chapter_heading_style_source = (
        template_defaults["chapter_heading_style"]
        if explicit_template
        else profile_defaults.get(
            "chapter_heading_style",
            compile_config.get("chapter_heading_style", template_defaults["chapter_heading_style"]),
        )
    )
    chapter_heading_style = str(chapter_heading_style_source).strip().lower()
    if chapter_heading_style not in {"title-only", "chapter-number-title", "section-number-title"}:
        raise CompileError(
            "chapter_heading_style must be 'title-only', 'chapter-number-title', or 'section-number-title'."
        )

    return CompileOptions(
        format_name=resolved_format,
        profile_name=resolved_profile,
        template_name=resolved_template,
        output_filename=str(compile_config.get("output_filename", "") or "").strip(),
        include_title_page=bool(
            template_defaults["include_title_page"]
            if explicit_template
            else profile_defaults.get(
                "include_title_page",
                compile_config.get("include_title_page", template_defaults["include_title_page"]),
            )
        ),
        include_part_headings=bool(
            template_defaults["include_part_headings"]
            if explicit_template
            else profile_defaults.get(
                "include_part_headings",
                compile_config.get("include_part_headings", template_defaults["include_part_headings"]),
            )
        ),
        chapter_heading_style=chapter_heading_style,
        backend_name=str(compile_config.get("backend", "auto") or "auto").strip().lower(),
        citation_style=str(research_compile.get("citation_style", "APA") or "APA").strip(),
        include_bibliography=bool(research_compile.get("include_bibliography", False)),
        include_reference_heading=bool(research_compile.get("include_reference_heading", True)),
        bibliography_title=str(research_compile.get("bibliography_title", "References") or "References").strip(),
    )


def _compiled_chapters(root: Path) -> list[ChapterDocument]:
    chapters = list_chapters(root)
    if not chapters:
        raise CompileError("No manuscript content exists yet. Create at least one chapter first.")

    populated_chapters = [chapter for chapter in chapters if chapter.body.strip()]
    if not populated_chapters:
        raise CompileError("No manuscript body text exists yet. Add text to at least one chapter first.")
    return populated_chapters


def _can_use_pandoc(options: CompileOptions) -> bool:
    if options.backend_name == "native":
        return False
    if options.backend_name == "pandoc":
        return True
    return shutil.which("pandoc") is not None


def _compile_project_with_pandoc(
    root: Path,
    project_title: str,
    author: str,
    chapters: list[ChapterDocument],
    target_path: Path,
    options: CompileOptions,
) -> None:
    pandoc_path = shutil.which("pandoc")
    if pandoc_path is None:
        raise CompileError("Pandoc was requested but is not installed or not available on PATH.")

    manuscript_text = _pandoc_markdown(root, project_title, author, chapters, options)
    with tempfile.TemporaryDirectory(prefix="openscribe-pandoc-") as temp_dir:
        temp_path = Path(temp_dir)
        source_path = temp_path / "manuscript.md"
        source_path.write_text(manuscript_text, encoding="utf-8")

        command = [
            pandoc_path,
            str(source_path),
            "--from",
            "markdown",
            "--to",
            options.format_name,
            "--output",
            str(target_path),
            "--standalone",
            "--metadata",
            f"title={project_title}",
        ]
        if author:
            command.extend(["--metadata", f"author={author}"])

        result = subprocess.run(command, capture_output=True, text=True, check=False)
        if result.returncode != 0:
            error_text = (result.stderr or result.stdout or "").strip()
            raise CompileError(f"Pandoc export failed. {error_text or 'Unknown Pandoc error.'}")


def _pandoc_markdown(
    root: Path,
    project_title: str,
    author: str,
    chapters: list[ChapterDocument],
    options: CompileOptions,
) -> str:
    lines: list[str] = [f"% {project_title}"]
    if author:
        lines.append(f"% {author}")
    lines.append("%")
    lines.append("")

    if options.include_title_page:
        lines.append(f"# {project_title}")
        lines.append("")
        if author:
            lines.append(author)
            lines.append("")

    current_part = None
    for chapter_number, chapter in enumerate(chapters, start=1):
        if chapter.part_id != current_part:
            current_part = chapter.part_id
            if options.include_part_headings:
                lines.append(f"# {chapter.part}")
                lines.append("")

        lines.append(f"## {_chapter_heading(chapter.title, chapter_number, options)}")
        lines.append("")
        for block in _paragraphs(chapter.body):
            lines.append(block)
            lines.append("")

    bibliography_sections = _bibliography_sections(root, options)
    if bibliography_sections:
        for section in bibliography_sections:
            lines.append(section)
            lines.append("")

    return "\n".join(lines).strip() + "\n"


def _compile_project_to_docx(
    root: Path,
    project_title: str,
    author: str,
    chapters: list[ChapterDocument],
    target_path: Path,
    options: CompileOptions,
) -> None:
    document = Document()
    document.core_properties.title = project_title
    if author:
        document.core_properties.author = author

    if options.include_title_page:
        document.add_heading(project_title, level=0)
        if author:
            document.add_paragraph(author)

    current_part = None
    for chapter_number, chapter in enumerate(chapters, start=1):
        if chapter.part_id != current_part:
            current_part = chapter.part_id
            if options.include_part_headings:
                document.add_page_break()
                document.add_heading(chapter.part, level=1)
        document.add_heading(_chapter_heading(chapter.title, chapter_number, options), level=2)
        for block in _paragraphs(chapter.body):
            document.add_paragraph(block)

    _append_bibliography_to_docx(document, root, options)
    document.save(target_path)


def _compile_project_to_pdf(
    root: Path,
    project_title: str,
    author: str,
    chapters: list[ChapterDocument],
    target_path: Path,
    options: CompileOptions,
) -> None:
    pdf = canvas.Canvas(str(target_path), pagesize=LETTER)
    width, height = LETTER
    left_margin = 72
    right_margin = width - 72
    top_margin = height - 72
    bottom_margin = 72
    line_height = 14
    y = top_margin

    pdf.setTitle(project_title)
    if author:
        pdf.setAuthor(author)

    def new_page() -> None:
        nonlocal y
        pdf.showPage()
        pdf.setFont("Times-Roman", 12)
        y = top_margin

    def ensure_space(lines_needed: int = 1) -> None:
        nonlocal y
        if y - (lines_needed * line_height) < bottom_margin:
            new_page()

    if options.include_title_page:
        pdf.setFont("Times-Bold", 20)
        ensure_space(2)
        pdf.drawString(left_margin, y, project_title)
        y -= line_height * 2

        if author:
            pdf.setFont("Times-Roman", 12)
            ensure_space(2)
            pdf.drawString(left_margin, y, author)
            y -= line_height * 2

    current_part = None
    for chapter_number, chapter in enumerate(chapters, start=1):
        if chapter.part_id != current_part:
            current_part = chapter.part_id
            if options.include_part_headings:
                new_page()
                pdf.setFont("Times-Bold", 16)
                ensure_space(2)
                pdf.drawString(left_margin, y, chapter.part)
                y -= line_height * 2

        pdf.setFont("Times-Bold", 14)
        ensure_space(2)
        pdf.drawString(left_margin, y, _chapter_heading(chapter.title, chapter_number, options))
        y -= line_height * 2

        pdf.setFont("Times-Roman", 12)
        for block in _paragraphs(chapter.body):
            wrapped_lines = _wrap_pdf_text(block, right_margin - left_margin, "Times-Roman", 12)
            ensure_space(len(wrapped_lines) + 1)
            for line in wrapped_lines:
                pdf.drawString(left_margin, y, line)
                y -= line_height
            y -= line_height

    bibliography_lines = _bibliography_render_lines(root, options)
    if bibliography_lines:
        pdf.setFont("Times-Bold", 14)
        ensure_space(2)
        pdf.drawString(left_margin, y, options.bibliography_title)
        y -= line_height * 2
        pdf.setFont("Times-Roman", 12)
        for entry in bibliography_lines:
            wrapped_lines = _wrap_pdf_text(entry, right_margin - left_margin, "Times-Roman", 12)
            ensure_space(len(wrapped_lines) + 1)
            for line in wrapped_lines:
                pdf.drawString(left_margin, y, line)
                y -= line_height
            y -= line_height

    pdf.save()


def _compile_project_to_epub(
    root: Path,
    project_title: str,
    author: str,
    chapters: list[ChapterDocument],
    target_path: Path,
    options: CompileOptions,
) -> None:
    book = epub.EpubBook()
    book.set_identifier(f"openscribe-{target_path.stem}")
    book.set_title(project_title)
    book.set_language("en")
    if author:
        book.add_author(author)

    title_page = None
    if options.include_title_page:
        title_page = epub.EpubHtml(title="Title Page", file_name="title.xhtml", lang="en")
        title_page.content = _epub_page(project_title, author, "")
        book.add_item(title_page)

    epub_items: list[epub.EpubHtml] = []
    current_part = None

    for chapter_index, chapter in enumerate(chapters, start=1):
        part_heading = ""
        if chapter.part_id != current_part:
            current_part = chapter.part_id
            if options.include_part_headings:
                part_heading = f"<h1>{escape(chapter.part)}</h1>"

        chapter_item = epub.EpubHtml(
            title=_chapter_heading(chapter.title, chapter_index, options),
            file_name=f"chapter-{chapter_index:02d}.xhtml",
            lang="en",
        )
        chapter_item.content = _epub_page(
            _chapter_heading(chapter.title, chapter_index, options),
            "",
            part_heading + _epub_body(chapter.body),
        )
        book.add_item(chapter_item)
        epub_items.append(chapter_item)

    bibliography_html = _epub_bibliography_page(root, options)
    if bibliography_html:
        bibliography_page = epub.EpubHtml(
            title=options.bibliography_title,
            file_name="references.xhtml",
            lang="en",
        )
        bibliography_page.content = bibliography_html
        book.add_item(bibliography_page)
        epub_items.append(bibliography_page)

    if title_page is not None:
        book.toc = tuple([title_page, *epub_items])
        book.spine = ["nav", title_page, *epub_items]
    else:
        book.toc = tuple(epub_items)
        book.spine = ["nav", *epub_items]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(str(target_path), book, {})


def _paragraphs(text: str) -> list[str]:
    visible_text = strip_scene_markers(text)
    return [block.strip() for block in visible_text.split("\n\n") if block.strip()]


def _wrap_pdf_text(text: str, max_width: float, font_name: str, font_size: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]

    lines: list[str] = []
    current_line = words[0]
    for word in words[1:]:
        candidate = f"{current_line} {word}"
        if stringWidth(candidate, font_name, font_size) <= max_width:
            current_line = candidate
            continue
        lines.append(current_line)
        current_line = word
    lines.append(current_line)
    return lines


def _epub_page(title: str, subtitle: str, body_html: str) -> str:
    subtitle_html = f"<p>{escape(subtitle)}</p>" if subtitle else ""
    return (
        "<html><head><title>"
        f"{escape(title)}"
        "</title></head><body>"
        f"<h1>{escape(title)}</h1>"
        f"{subtitle_html}"
        f"{body_html}"
        "</body></html>"
    )


def _epub_body(text: str) -> str:
    blocks = [f"<p>{escape(block)}</p>" for block in _paragraphs(text)]
    return "".join(blocks)


def _bibliography_sections(root: Path, options: CompileOptions) -> list[str]:
    entries = _bibliography_render_lines(root, options)
    if not entries:
        return []
    if options.include_reference_heading:
        return [options.bibliography_title, "\n".join(entries)]
    return ["\n".join(entries)]


def _bibliography_render_lines(root: Path, options: CompileOptions) -> list[str]:
    if not options.include_bibliography:
        return []
    entries = [_source_entry(source, options.citation_style) for source in list_source_notes(root)]
    return [entry for entry in entries if entry.strip()]


def _append_bibliography_to_docx(document: Document, root: Path, options: CompileOptions) -> None:
    entries = _bibliography_render_lines(root, options)
    if not entries:
        return
    if options.include_reference_heading:
        document.add_page_break()
        document.add_heading(options.bibliography_title, level=1)
    for entry in entries:
        document.add_paragraph(entry)


def _epub_bibliography_page(root: Path, options: CompileOptions) -> str:
    entries = _bibliography_render_lines(root, options)
    if not entries:
        return ""
    body = "".join(f"<p>{escape(entry)}</p>" for entry in entries)
    return _epub_page(options.bibliography_title, "", body)


def _source_entry(source, citation_style: str) -> str:
    author = source.author.strip()
    year = source.year.strip()
    url = source.url.strip()
    title = source.title.strip()
    style = citation_style.strip().upper()
    if style == "MLA":
        parts = [f"{author}." if author else "", f"{title}.", year, url]
        return " ".join(part for part in parts if part).strip()
    if style == "CHICAGO":
        parts = [author, f'"{title}."' if title else "", year, url]
        return ". ".join(part.strip().rstrip(".") for part in parts if part).strip() + "."
    parts = [f"{author} ({year})." if author and year else author or f"({year})." if year else "", title + ".", url]
    return " ".join(part for part in parts if part).strip()


def _chapter_heading(title: str, chapter_number: int, options: CompileOptions) -> str:
    if options.chapter_heading_style == "section-number-title":
        return f"{chapter_number}. {title}"
    if options.chapter_heading_style == "chapter-number-title":
        return f"Chapter {chapter_number}: {title}"
    return title
